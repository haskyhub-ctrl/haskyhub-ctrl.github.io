"""Bảng theo dõi realtime của admin: số cơ sở đã tham gia khảo sát theo đội quản lý.

Một cơ sở được tính là "đã tham gia" khi có ít nhất một bài khảo sát đã hoàn thành. Chỉ tính các tài khoản
cơ sở nạp từ danh sách quản lý (có đội quản lý), không tính người dân khảo sát tự do.
"""
import io
from datetime import datetime, timedelta
from typing import Optional

from fastapi import APIRouter, Depends, Query
from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import case, func
from sqlalchemy.orm import Session

from database import get_db
from middleware.auth_middleware import get_current_user
from middleware.rbac import require_role
from models import Assessment, User

router = APIRouter(prefix="/api/admin/teams", tags=["Admin Teams"])

RISK_LABELS = {"low": "Thấp", "safe": "Thấp", "medium": "Trung bình", "high": "Cao", "critical": "Rất cao"}

GROUP_LABELS = {"pc07": "Đội PC07", "commune": "Công an xã, phường"}
SORTS = {
    "percent": lambda t: (-t["percent"], -t["participated"], t["team"]),
    "participated": lambda t: (-t["participated"], -t["percent"], t["team"]),
    "total": lambda t: (-t["total"], -t["percent"], t["team"]),
    "name": lambda t: (t["team"],),
}


def _group_of(team: str) -> str:
    """Đội thuộc PC07 (Đội, Tổ...) hay Công an xã, phường (tên đơn vị bắt đầu bằng UBND)."""
    return "commune" if (team or "").strip().upper().startswith("UBND") else "pc07"


def _apply_team_filters(qry, group, team, q):
    if group == "commune":
        qry = qry.filter(User.management_team.ilike("UBND%"))
    elif group == "pc07":
        qry = qry.filter(~User.management_team.ilike("UBND%"))
    if team:
        qry = qry.filter(User.management_team == team)
    if q:
        qry = qry.filter(User.management_team.ilike(f"%{q.strip()}%"))
    return qry


def _latest_completed(db: Session):
    """Mỗi cơ sở: bài khảo sát hoàn thành gần nhất (mức nguy cơ, thời điểm) và số bài đã hoàn thành."""
    ranked = (
        db.query(
            Assessment.user_id.label("uid"),
            Assessment.id.label("aid"),
            Assessment.risk_level.label("risk"),
            Assessment.risk_percentage.label("pct"),
            Assessment.completed_at.label("done_at"),
            case((Assessment.mitigation_data.like('%"resolved": true%'), 1), else_=0).label("mit"),
            func.row_number().over(partition_by=Assessment.user_id, order_by=Assessment.completed_at.desc()).label("rn"),
            func.count(Assessment.id).over(partition_by=Assessment.user_id).label("n"),
        )
        .filter(Assessment.status == "completed")
        .subquery()
    )
    return db.query(ranked).filter(ranked.c.rn == 1).subquery()


def _iso(dt):
    return dt.isoformat() if dt else None


def _mit_expr(latest):
    """Số cơ sở đã cập nhật tiến độ khắc phục (có ít nhất một hạng mục tích đã khắc phục)."""
    return func.coalesce(func.sum(latest.c.mit), 0)


def _high_expr(latest):
    return func.coalesce(func.sum(case((latest.c.risk.in_(("high", "critical")), 1), else_=0)), 0)


def _overview_data(db, ward, q, group, team, sort):
    latest = _latest_completed(db)

    base = (
        db.query(
            User.management_team.label("team"),
            func.count(User.id).label("total"),
            func.count(latest.c.uid).label("participated"),
            func.max(latest.c.done_at).label("last_at"),
            _high_expr(latest).label("high"),
            _mit_expr(latest).label("mit"),
        )
        .outerjoin(latest, latest.c.uid == User.id)
        .filter(User.management_team.isnot(None))
    )
    if ward:
        base = base.filter(User.ward == ward)
    base = _apply_team_filters(base, group, team, q)
    rows = base.group_by(User.management_team).all()

    teams = []
    for r in rows:
        total, part = int(r.total), int(r.participated)
        teams.append({
            "team": r.team,
            "total": total,
            "participated": part,
            "remaining": total - part,
            "percent": round(part * 100.0 / total, 1) if total else 0.0,
            "high_risk": int(r.high or 0),
            "mitigated": int(r.mit or 0),
            "last_survey_at": _iso(r.last_at),
            "group": _group_of(r.team),
            "group_label": GROUP_LABELS[_group_of(r.team)],
        })
    teams.sort(key=SORTS.get(sort, SORTS["percent"]))

    # Thống kê theo xã phường (tôn trọng bộ lọc tên đội, không áp bộ lọc xã phường để còn chọn được)
    wq = (
        db.query(
            User.ward.label("ward"),
            func.count(User.id).label("total"),
            func.count(latest.c.uid).label("participated"),
            _mit_expr(latest).label("mit"),
        )
        .outerjoin(latest, latest.c.uid == User.id)
        .filter(User.management_team.isnot(None))
    )
    wq = _apply_team_filters(wq, group, team, q)
    wards = [
        {"ward": w.ward or "(chưa có)", "value": w.ward or "", "total": int(w.total), "participated": int(w.participated),
         "mitigated": int(w.mit or 0),
         "percent": round(int(w.participated) * 100.0 / int(w.total), 1) if w.total else 0.0}
        for w in wq.group_by(User.ward).all()
    ]
    wards.sort(key=lambda w: (-w["percent"], -w["participated"], w["ward"]))

    total = sum(t["total"] for t in teams)
    part = sum(t["participated"] for t in teams)
    since = datetime.utcnow() - timedelta(hours=24)
    l24 = (
        db.query(func.count(func.distinct(Assessment.user_id)))
        .join(User, User.id == Assessment.user_id)
        .filter(Assessment.status == "completed", Assessment.completed_at >= since, User.management_team.isnot(None))
    )
    if ward:
        l24 = l24.filter(User.ward == ward)
    l24 = _apply_team_filters(l24, group, team, q)
    last_24h = l24.scalar() or 0
    return {
        "generated_at": datetime.utcnow().isoformat() + "Z",
        "totals": {
            "facilities": total,
            "participated": part,
            "remaining": total - part,
            "percent": round(part * 100.0 / total, 1) if total else 0.0,
            "high_risk": sum(t["high_risk"] for t in teams),
            "mitigated": sum(t["mitigated"] for t in teams),
            "last_24h": int(last_24h),
            "teams": len(teams),
        },
        "teams": teams,
        "wards": wards,
        "filters": {"ward": ward or "", "q": q or "", "group": group or "", "team": team or "", "sort": sort},
    }


@router.get("/overview")
def teams_overview(
    ward: Optional[str] = Query(None, description="Lọc theo xã phường"),
    q: Optional[str] = Query(None, description="Tìm theo tên đội"),
    group: Optional[str] = Query(None, pattern="^(pc07|commune)$", description="pc07 hoặc commune (Công an xã, phường)"),
    team: Optional[str] = Query(None, description="Chọn đúng một đội"),
    sort: str = Query("percent", pattern="^(percent|participated|total|name)$"),
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    """Tổng hợp theo đội quản lý: số cơ sở, số đã tham gia, tỷ lệ, bài mới nhất (mặc định xếp theo tỷ lệ cao nhất)."""
    require_role("admin", "superadmin")(current_user)
    return _overview_data(db, ward, q, group, team, sort)


def _style_sheet(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1A3A6B")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"


def _xlsx_response(wb, filename):
    buf = io.BytesIO()
    wb.save(buf)
    from urllib.parse import quote
    return Response(
        content=buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}"},
    )


@router.get("/overview/export")
def overview_export(
    ward: Optional[str] = None,
    q: Optional[str] = None,
    group: Optional[str] = Query(None, pattern="^(pc07|commune)$"),
    team: Optional[str] = None,
    sort: str = Query("percent", pattern="^(percent|participated|total|name)$"),
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    """Xuất Excel bảng theo dõi theo các bộ lọc đang chọn (sheet theo đội và sheet theo xã phường)."""
    require_role("admin", "superadmin")(current_user)
    d = _overview_data(db, ward, q, group, team, sort)
    wb = Workbook()
    ws = wb.active
    ws.title = "Theo đội"
    ws.append(["STT", "Đội quản lý", "Nhóm", "Tổng cơ sở", "Đã khảo sát", "Chưa khảo sát", "Tỷ lệ tham gia (%)", "Nguy cơ cao, rất cao", "Đã cập nhật khắc phục", "Khảo sát gần nhất"])
    for i, t in enumerate(d["teams"], 1):
        ws.append([i, t["team"], t["group_label"], t["total"], t["participated"], t["remaining"], t["percent"], t["high_risk"], t["mitigated"],
                   (datetime.fromisoformat(t["last_survey_at"]).strftime("%d/%m/%Y %H:%M") if t["last_survey_at"] else "")])
    tt = d["totals"]
    ws.append(["", "TỔNG", "", tt["facilities"], tt["participated"], tt["remaining"], tt["percent"], tt["high_risk"], tt["mitigated"], ""])
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    _style_sheet(ws, [6, 52, 20, 12, 12, 14, 16, 16, 18, 20])
    ws2 = wb.create_sheet("Theo xã phường")
    ws2.append(["STT", "Xã phường", "Tổng cơ sở", "Đã khảo sát", "Đã cập nhật khắc phục", "Tỷ lệ tham gia (%)"])
    for i, w in enumerate(d["wards"], 1):
        ws2.append([i, w["ward"], w["total"], w["participated"], w["mitigated"], w["percent"]])
    _style_sheet(ws2, [6, 34, 12, 12, 18, 16])
    return _xlsx_response(wb, f"Theo_doi_khao_sat_{datetime.now():%Y%m%d_%H%M}.xlsx")


@router.get("/map-wards")
def map_wards(
    group: Optional[str] = Query(None, pattern="^(pc07|commune)$"),
    team: Optional[str] = None,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    """Thống kê theo xã phường để vẽ bản đồ quản lý: tổng cơ sở, đã khảo sát, nguy cơ trung bình, nguy cơ cao, tâm tọa độ."""
    require_role("admin", "superadmin")(current_user)
    latest = _latest_completed(db)
    qry = (
        db.query(
            User.ward,
            func.count(User.id),
            func.count(latest.c.uid),
            func.avg(latest.c.pct),
            _high_expr(latest),
            _mit_expr(latest),
            func.avg(User.latitude),
            func.avg(User.longitude),
        )
        .outerjoin(latest, latest.c.uid == User.id)
        .filter(User.management_team.isnot(None))
    )
    qry = _apply_team_filters(qry, group, team, None).group_by(User.ward)
    out = []
    for ward, total, surveyed, avg_pct, high, mit, lat, lng in qry.all():
        if lat is None or lng is None or not ward:
            continue
        total, surveyed = int(total), int(surveyed)
        out.append({
            "ward": ward, "total": total, "surveyed": surveyed,
            "coverage": round(surveyed * 100.0 / total, 1) if total else 0.0,
            "avg_risk": round(float(avg_pct), 1) if avg_pct is not None else None,
            "high": int(high or 0),
            "high_rate": round(int(high or 0) * 100.0 / surveyed, 1) if surveyed else None,
            "mitigated": int(mit or 0),
            "lat": round(float(lat), 6), "lng": round(float(lng), 6),
        })
    return {"generated_at": datetime.utcnow().isoformat() + "Z", "wards": out}


def _detail_query(db, latest, team, ward, status, q):
    qry = (
        db.query(
            User.id, User.facility_code, User.full_name, User.ward, User.facility_type_label, User.facility_types,
            latest.c.aid, latest.c.risk, latest.c.pct, latest.c.done_at, latest.c.n, latest.c.mit,
        )
        .outerjoin(latest, latest.c.uid == User.id)
        .filter(User.management_team == team)
    )
    if ward:
        qry = qry.filter(User.ward == ward)
    if status == "done":
        qry = qry.filter(latest.c.uid.isnot(None))
    elif status == "todo":
        qry = qry.filter(latest.c.uid.is_(None))
    elif status == "mitigated":
        qry = qry.filter(latest.c.mit == 1)
    if q:
        like = f"%{q.strip()}%"
        qry = qry.filter(User.full_name.ilike(like) | User.facility_code.like(like))
    return qry


@router.get("/detail")
def team_detail(
    team: str = Query(..., description="Tên đội quản lý"),
    ward: Optional[str] = None,
    status: str = Query("all", pattern="^(all|done|todo|mitigated)$"),
    q: Optional[str] = None,
    page: int = Query(1, ge=1),
    page_size: int = Query(50, ge=10, le=200),
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    """Danh sách cơ sở của một đội, kèm trạng thái đã hoặc chưa khảo sát. Lọc theo xã phường, trạng thái, từ khóa."""
    require_role("admin", "superadmin")(current_user)
    latest = _latest_completed(db)

    total_items = _detail_query(db, latest, team, ward, status, q).count()
    qry = _detail_query(db, latest, team, ward, status, q)
    if status == "done":
        qry = qry.order_by(latest.c.done_at.desc())
    else:
        qry = qry.order_by(User.ward, User.full_name)
    rows = qry.offset((page - 1) * page_size).limit(page_size).all()

    # Tổng hợp của đội (áp bộ lọc xã phường) và thống kê từng xã phường để chọn lọc
    summary_q = (
        db.query(func.count(User.id), func.count(latest.c.uid), _mit_expr(latest))
        .outerjoin(latest, latest.c.uid == User.id)
        .filter(User.management_team == team)
    )
    if ward:
        summary_q = summary_q.filter(User.ward == ward)
    s_total, s_part, s_mit = summary_q.one()
    wards = [
        {"ward": w or "(chưa có)", "value": w or "", "total": int(t), "participated": int(p)}
        for w, t, p in (
            db.query(User.ward, func.count(User.id), func.count(latest.c.uid))
            .outerjoin(latest, latest.c.uid == User.id)
            .filter(User.management_team == team)
            .group_by(User.ward)
            .order_by(User.ward)
            .all()
        )
    ]
    return {
        "generated_at": datetime.utcnow().isoformat() + "Z",
        "team": team,
        "summary": {
            "total": int(s_total), "participated": int(s_part), "remaining": int(s_total) - int(s_part),
            "mitigated": int(s_mit or 0),
            "percent": round(int(s_part) * 100.0 / int(s_total), 1) if s_total else 0.0,
        },
        "wards": wards,
        "page": page,
        "page_size": page_size,
        "total_items": total_items,
        "items": [
            {
                "user_id": r.id,
                "code": r.facility_code,
                "name": r.full_name,
                "ward": r.ward,
                "type_label": r.facility_type_label or "",
                "participated": r.aid is not None,
                "assessment_id": r.aid,
                "risk_level": r.risk,
                "risk_label": RISK_LABELS.get(r.risk, ""),
                "risk_percentage": r.pct,
                "completed_at": _iso(r.done_at),
                "surveys": int(r.n) if r.n else 0,
                "mitigated": bool(r.mit),
            }
            for r in rows
        ],
    }


@router.get("/export")
def team_export(
    team: str = Query(...),
    ward: Optional[str] = None,
    status: str = Query("done", pattern="^(all|done|todo|mitigated)$"),
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    """Tải danh sách cơ sở của đội dạng Excel (.xlsx)."""
    require_role("admin", "superadmin")(current_user)
    latest = _latest_completed(db)
    rows = _detail_query(db, latest, team, ward, status, None).order_by(User.ward, User.full_name).all()
    wb = Workbook()
    ws = wb.active
    ws.title = "Danh sách cơ sở"
    ws.append(["STT", "Mã cơ sở", "Tên cơ sở", "Xã phường", "Loại hình", "Trạng thái", "Mức nguy cơ", "Tỷ lệ nguy cơ (%)", "Thời điểm khảo sát", "Số bài", "Đã cập nhật khắc phục"])
    for i, r in enumerate(rows, 1):
        ws.append([
            i, str(r.facility_code or ""), r.full_name, r.ward or "", r.facility_type_label or "",
            "Đã khảo sát" if r.aid else "Chưa khảo sát",
            RISK_LABELS.get(r.risk, ""), r.pct if r.pct is not None else "",
            r.done_at.strftime("%d/%m/%Y %H:%M") if r.done_at else "", int(r.n) if r.n else 0, "Có" if r.mit else "",
        ])
    _style_sheet(ws, [6, 18, 48, 26, 40, 16, 14, 16, 20, 8, 18])
    safe = "".join(ch if ch.isalnum() or ch in " _-" else "_" for ch in team).strip().replace(" ", "_")
    return _xlsx_response(wb, f"DS_{safe}_{datetime.now():%Y%m%d}.xlsx")


# ------------------------------------------------------------------ rà soát tọa độ
def _audit_rows(db, team=None):
    from utils import geo_wards
    qry = db.query(User.facility_code, User.full_name, User.management_team, User.ward, User.latitude, User.longitude) \
        .filter(User.management_team.isnot(None))
    if team:
        qry = qry.filter(User.management_team == team)
    rows = qry.all()
    have = [r for r in rows if r.latitude is not None and r.longitude is not None]
    actual = geo_wards.locate_many([(r.latitude, r.longitude) for r in have])
    out = []
    for r, a in zip(have, actual):
        if a is None:
            status = "outside"
        elif geo_wards.ward_key(a) == geo_wards.ward_key(r.ward):
            status = "ok"
        else:
            status = "mismatch"
        out.append((r, a, status))
    return rows, out


@router.get("/geo-audit")
def geo_audit(
    team: Optional[str] = None,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    """Rà tọa độ: so xã phường ghi trong danh sách với xã phường thực tế chứa điểm tọa độ (theo ranh giới địa giới)."""
    require_role("admin", "superadmin")(current_user)
    rows, out = _audit_rows(db, team)
    summary = {"total": len(rows), "no_coords": len(rows) - len(out), "ok": 0, "mismatch": 0, "outside": 0}
    by_team = {}
    for r, a, st in out:
        summary[st] += 1
        t = by_team.setdefault(r.management_team, {"team": r.management_team, "checked": 0, "mismatch": 0, "outside": 0})
        t["checked"] += 1
        if st != "ok":
            t[st] += 1
    teams_list = sorted(by_team.values(), key=lambda t: -(t["mismatch"] + t["outside"]))
    for t in teams_list:
        bad = t["mismatch"] + t["outside"]
        t["bad"] = bad
        t["bad_percent"] = round(bad * 100.0 / t["checked"], 1) if t["checked"] else 0.0
    return {"summary": summary, "teams": teams_list[:60]}


@router.get("/geo-audit/export")
def geo_audit_export(
    team: Optional[str] = None,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    """Tải danh sách cơ sở có tọa độ nghi sai hoặc chưa có tọa độ (Excel), kèm đường dẫn mở trên Google Maps."""
    require_role("admin", "superadmin")(current_user)
    rows, out = _audit_rows(db, team)
    wb = Workbook()
    ws = wb.active
    ws.title = "Tọa độ cần kiểm tra"
    ws.append(["STT", "Mã cơ sở", "Tên cơ sở", "Đội quản lý", "Xã phường ghi trong danh sách", "Xã phường theo tọa độ", "Kết luận", "Vĩ độ", "Kinh độ", "Mở bản đồ"])
    n = 0
    for r, a, st in out:
        if st == "ok":
            continue
        n += 1
        ws.append([n, str(r.facility_code or ""), r.full_name, r.management_team, r.ward or "", a or "(ngoài địa giới tỉnh)",
                   "Lệch xã phường" if st == "mismatch" else "Ngoài địa giới tỉnh", r.latitude, r.longitude,
                   f"https://www.google.com/maps?q={r.latitude},{r.longitude}"])
    have = {(x[0].facility_code) for x in out}
    for r in rows:
        if r.facility_code not in have:
            n += 1
            ws.append([n, str(r.facility_code or ""), r.full_name, r.management_team, r.ward or "", "", "Chưa có tọa độ", "", "", ""])
    _style_sheet(ws, [6, 18, 44, 40, 24, 24, 20, 12, 12, 46])
    return _xlsx_response(wb, f"Ra_soat_toa_do_{datetime.now():%Y%m%d}.xlsx")


# ------------------------------------------------------------------ tìm cơ sở theo tên để xem vị trí
_facility_cache = {"at": None, "rows": []}


def _strip(text):
    import unicodedata
    t = unicodedata.normalize("NFD", (text or "").lower().replace("đ", "d"))
    return "".join(c for c in t if unicodedata.category(c) != "Mn")


def _facility_index(db):
    now = datetime.utcnow()
    if _facility_cache["at"] is None or (now - _facility_cache["at"]).total_seconds() > 300:
        rows = db.query(User.id, User.facility_code, User.full_name, User.management_team, User.ward, User.latitude, User.longitude) \
            .filter(User.management_team.isnot(None), User.latitude.isnot(None), User.longitude.isnot(None)).all()
        _facility_cache["rows"] = [(r, _strip(r.full_name)) for r in rows]
        _facility_cache["at"] = now
    return _facility_cache["rows"]


@router.get("/facility-search")
def facility_search(
    q: str = Query(..., min_length=2, max_length=80),
    limit: int = Query(15, ge=1, le=30),
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    """Tìm cơ sở theo tên (không cần gõ dấu) hoặc mã cơ sở, trả về vị trí để mở trên bản đồ."""
    require_role("admin", "superadmin")(current_user)
    q = q.strip()
    tokens = [t for t in _strip(q).split() if t]
    hits = []
    for r, name in _facility_index(db):
        if q.isdigit():
            if (r.facility_code or "").startswith(q):
                hits.append((0, r))
        elif all(t in name for t in tokens):
            hits.append((0 if name.startswith(tokens[0]) else 1, r))
        if len(hits) >= 400:
            break
    hits.sort(key=lambda h: (h[0], len(h[1].full_name or "")))
    ids = [r.id for _, r in hits[:limit]]
    latest = _latest_completed(db)
    done = {uid: (risk, pct) for uid, risk, pct in db.query(latest.c.uid, latest.c.risk, latest.c.pct).filter(latest.c.uid.in_(ids)).all()} if ids else {}
    return [
        {
            "code": r.facility_code, "name": r.full_name, "team": r.management_team, "ward": r.ward,
            "lat": r.latitude, "lng": r.longitude,
            "surveyed": r.id in done, "risk_level": done.get(r.id, (None, None))[0],
            "risk_label": RISK_LABELS.get(done.get(r.id, (None, None))[0], ""),
        }
        for _, r in hits[:limit]
    ]
